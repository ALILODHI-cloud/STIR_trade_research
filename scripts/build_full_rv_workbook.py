#!/usr/bin/env python3
"""Build the exhaustive three-curve STIR relative-value workbook."""
from __future__ import annotations

import argparse
import shutil
from itertools import permutations
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
PANEL = ROOT / "data" / "cache" / "stir_curves" / "panel.csv"
OUT_DIR = ROOT / "research"
LATEST = OUT_DIR / "stir_full_rv_universe_latest.xlsx"
SAMPLE_START = pd.Timestamp("2026-03-01")

MARKET_ORDER = {"SOFR": 0, "EURIBOR": 1, "SONIA": 2}
MONTH_ORDER = {
    "F": 1,
    "G": 2,
    "H": 3,
    "J": 4,
    "K": 5,
    "M": 6,
    "N": 7,
    "Q": 8,
    "U": 9,
    "V": 10,
    "X": 11,
    "Z": 12,
}
MONTH_NAME = {
    "F": "January",
    "G": "February",
    "H": "March",
    "J": "April",
    "K": "May",
    "M": "June",
    "N": "July",
    "Q": "August",
    "U": "September",
    "V": "October",
    "X": "November",
    "Z": "December",
}
ROOT_SYMBOL = {"SOFR": "SQ", "EURIBOR": "IM", "SONIA": "J8"}

SEGMENTS = {
    "BUILDUP": {
        "label": "September 2027 minus December 2026",
        "weights": {"Z26": -1, "U27": 1},
    },
    "REVERSAL": {
        "label": "December 2028 minus September 2027",
        "weights": {"U27": -1, "Z28": 1},
    },
    "JUN_CAL": {
        "label": "June 2028 minus June 2027",
        "weights": {"M27": -1, "M28": 1},
    },
    "SEP_CAL": {
        "label": "September 2028 minus September 2027",
        "weights": {"U27": -1, "U28": 1},
    },
    "DEC_CAL": {
        "label": "December 2028 minus December 2027",
        "weights": {"Z27": -1, "Z28": 1},
    },
    "PEAK_CURV": {
        "label": "2× September 2027 minus December 2026 and December 2028",
        "weights": {"Z26": -1, "U27": 2, "Z28": -1},
    },
    "FRONT_BACK": {
        "label": "December 2028 minus December 2026",
        "weights": {"Z26": -1, "Z28": 1},
    },
    "MAR27_BUILD": {
        "label": "March 2027 minus December 2026",
        "weights": {"Z26": -1, "H27": 1},
    },
    "JUN27_BUILD": {
        "label": "June 2027 minus December 2026",
        "weights": {"Z26": -1, "M27": 1},
    },
}

NAVY = "14213D"
WHITE = "FFFFFF"
PALE = "F7F9FC"
RED = "F4CCCC"
AMBER = "FCE5CD"
THIN = Side(style="thin", color="D9E1F2")


def month_label(contract: str) -> str:
    return f"{MONTH_NAME[contract[0]]} 20{contract[1:]}"


def contract_key(contract: str) -> tuple[int, int]:
    return 2000 + int(contract[1:]), MONTH_ORDER[contract[0]]


def safe_beta(y: pd.Series, x: pd.Series) -> float:
    a = pd.concat([y, x], axis=1).dropna().tail(30)
    if len(a) < 15 or float(a.iloc[:, 1].var()) == 0:
        return np.nan
    return float(a.iloc[:, 0].cov(a.iloc[:, 1]) / a.iloc[:, 1].var())


def series_stats(level: pd.Series) -> dict:
    s = level.replace([np.inf, -np.inf], np.nan).dropna()
    current = float(s.iloc[-1])
    std = float(s.std(ddof=1))
    percentile = float((s <= current).mean() * 100)
    return {
        "observations": len(s),
        "current": current,
        "previous": float(s.iloc[-2]),
        "change_1d": current - float(s.iloc[-2]),
        "change_5d": current - float(s.iloc[-6]) if len(s) > 5 else np.nan,
        "change_20d": current - float(s.iloc[-21]) if len(s) > 20 else np.nan,
        "mean": float(s.mean()),
        "median": float(s.median()),
        "std": std,
        "zscore": (current - float(s.mean())) / std if std else np.nan,
        "minimum": float(s.min()),
        "min_date": str(s.idxmin().date()),
        "maximum": float(s.max()),
        "max_date": str(s.idxmax().date()),
        "percentile": percentile,
        "tail_percentile": min(percentile, 100 - percentile),
    }


def beta_diagnostics(long_level: pd.Series, short_level: pd.Series, pnl_sign: int) -> dict:
    """pnl_sign=-1 for receive-rate contract pair; +1 for long segment pair."""
    changes = pd.concat(
        [long_level.diff().rename("long"), short_level.diff().rename("short")],
        axis=1,
    ).dropna()
    last = changes.tail(30)
    beta = safe_beta(changes.long, changes.short)
    corr = float(last.long.corr(last.short)) if len(last) else np.nan
    pnl_1to1 = pnl_sign * (last.long - last.short)
    pnl_beta = pnl_sign * (last.long - beta * last.short)
    if np.isnan(beta):
        beta_bucket = "insufficient"
    elif beta < 0.8:
        beta_bucket = "<0.8: one short unit overhedges long leg"
    elif beta <= 1.2:
        beta_bucket = "0.8–1.2: broadly matched"
    else:
        beta_bucket = ">1.2: one short unit underhedges long leg"
    return {
        "beta30_long_on_short": beta,
        "beta_lt_1": bool(beta < 1) if not np.isnan(beta) else None,
        "corr30": corr,
        "r_squared30": corr**2 if not np.isnan(corr) else np.nan,
        "long_vol30": float(last.long.std(ddof=1)),
        "short_vol30": float(last.short.std(ddof=1)),
        "residual_vol30": float((last.long - beta * last.short).std(ddof=1)),
        "pnl_1to1_vol30": float(pnl_1to1.std(ddof=1)),
        "pnl_betahedged_vol30": float(pnl_beta.std(ddof=1)),
        "pnl_1to1_mean30": float(pnl_1to1.mean()),
        "pnl_1to1_win_pct30": float((pnl_1to1 > 0).mean() * 100),
        "beta_interpretation": beta_bucket,
    }


def load_panel(asof: str | None = None):
    panel = pd.read_csv(PANEL, parse_dates=["date"])
    end = pd.Timestamp(asof) if asof else panel.date.max()
    panel = panel[panel.date <= end].copy()
    end = panel.date.max()
    active = panel[panel.date == end][["curve", "symbol", "contract"]].drop_duplicates()
    active["_market"] = active.curve.map(MARKET_ORDER)
    active["_contract"] = active.contract.map(contract_key)
    active = active.sort_values(["_market", "_contract"]).drop(
        columns=["_market", "_contract"]
    )
    ids = {
        (r.curve, r.contract): f"{r.curve}_{r.contract}"
        for r in active.itertuples(index=False)
    }
    sample = panel[
        (panel.date >= SAMPLE_START)
        & panel.set_index(["curve", "contract"]).index.isin(ids.keys())
    ].copy()
    rates = sample.pivot_table(
        index="date", columns=["curve", "contract"], values="rate"
    ).sort_index()
    prices = sample.pivot_table(
        index="date", columns=["curve", "contract"], values="price"
    ).sort_index()
    rates.columns = [ids[c] for c in rates.columns]
    prices.columns = [ids[c] for c in prices.columns]
    return panel, active, rates, prices, end


def contract_directory(active: pd.DataFrame, rates: pd.DataFrame, prices: pd.DataFrame):
    rows = []
    for r in active.itertuples(index=False):
        cid = f"{r.curve}_{r.contract}"
        s = rates[cid].dropna()
        rows.append(
            {
                "contract_id": cid,
                "market": r.curve,
                "contract_month": month_label(r.contract),
                "vendor_symbol": r.symbol,
                "current_price": float(prices[cid].dropna().iloc[-1]),
                "current_rate": float(s.iloc[-1]),
                "rate_move_1d_bp": float((s.iloc[-1] - s.iloc[-2]) * 100),
                "rate_move_5d_bp": float((s.iloc[-1] - s.iloc[-6]) * 100)
                if len(s) > 5
                else np.nan,
            }
        )
    return pd.DataFrame(rows)


def contract_pairs(directory: pd.DataFrame, rates: pd.DataFrame) -> pd.DataFrame:
    meta = directory.set_index("contract_id").to_dict("index")
    rows = []
    for long_id, short_id in permutations(directory.contract_id, 2):
        long_s = rates[long_id] * 100
        short_s = rates[short_id] * 100
        aligned = pd.concat([long_s, short_s], axis=1).dropna()
        if len(aligned) < 31:
            continue
        long_s, short_s = aligned.iloc[:, 0], aligned.iloc[:, 1]
        spread = long_s - short_s
        st = series_stats(spread)
        diag = beta_diagnostics(long_s, short_s, pnl_sign=-1)
        lm, sm = meta[long_id], meta[short_id]
        if lm["market"] == sm["market"]:
            pair_type = "within market"
        elif lm["contract_month"] == sm["contract_month"]:
            pair_type = "cross market, same month"
        else:
            pair_type = "cross market, different month"
        rows.append(
            {
                "long_receive_leg": long_id,
                "long_month": lm["contract_month"],
                "short_pay_leg": short_id,
                "short_month": sm["contract_month"],
                "pair_type": pair_type,
                "rate_spread_long_minus_short_bp": st["current"],
                "futures_price_spread_long_minus_short": -st["current"] / 100,
                "spread_change_1d_bp": st["change_1d"],
                "trade_pnl_1d_bp": -st["change_1d"],
                "trade_pnl_5d_bp": -st["change_5d"],
                "spread_mean_bp": st["mean"],
                "spread_zscore": st["zscore"],
                "spread_percentile": st["percentile"],
                "spread_min_bp": st["minimum"],
                "spread_max_bp": st["maximum"],
                "entry_read": (
                    "rich for receive-long/pay-short mean reversion"
                    if st["percentile"] >= 95
                    else "cheap; inverse direction has better mean-reversion entry"
                    if st["percentile"] <= 5
                    else "not a tail valuation"
                ),
                **diag,
            }
        )
    return pd.DataFrame(rows)


def segment_data(rates: pd.DataFrame):
    series = {}
    directory = []
    for market in MARKET_ORDER:
        for segment, spec in SEGMENTS.items():
            needed = [f"{market}_{c}" for c in spec["weights"]]
            if not all(c in rates for c in needed):
                continue
            sid = f"{market}_{segment}"
            value = sum(
                weight * rates[f"{market}_{contract}"] * 100
                for contract, weight in spec["weights"].items()
            )
            value = value.dropna()
            series[sid] = value
            st = series_stats(value)
            directory.append(
                {
                    "segment_id": sid,
                    "market": market,
                    "segment": segment,
                    "definition": spec["label"],
                    "contract_rate_weights": ", ".join(
                        f"{month_label(c)} {w:+g}" for c, w in spec["weights"].items()
                    ),
                    "current_bp": st["current"],
                    "change_1d_bp": st["change_1d"],
                    "change_5d_bp": st["change_5d"],
                    "mean_bp": st["mean"],
                    "zscore": st["zscore"],
                    "percentile": st["percentile"],
                    "min_bp": st["minimum"],
                    "max_bp": st["maximum"],
                }
            )
    wide = pd.concat(series, axis=1, sort=True)
    return pd.DataFrame(directory), wide


def segment_weights(segment_id: str, multiplier: float) -> dict[tuple[str, str], float]:
    market, segment = segment_id.split("_", 1)
    return {
        (market, contract): multiplier * weight
        for contract, weight in SEGMENTS[segment]["weights"].items()
    }


def trade_legs(long_id: str, short_id: str, beta: float = 1.0) -> str:
    weights: dict[tuple[str, str], float] = {}
    # Long segment P&L = +dSegment; short segment P&L = -beta*dSegment.
    for key, value in segment_weights(long_id, 1.0).items():
        weights[key] = weights.get(key, 0) + value
    for key, value in segment_weights(short_id, -beta).items():
        weights[key] = weights.get(key, 0) + value
    legs = []
    for (market, contract), weight in sorted(
        weights.items(), key=lambda x: (MARKET_ORDER[x[0][0]], contract_key(x[0][1]))
    ):
        if abs(weight) < 1e-10:
            continue
        action = "PAY" if weight > 0 else "RECEIVE"
        size = abs(weight)
        size_text = "" if np.isclose(size, 1) else f"{size:.2f}× "
        legs.append(f"{action} {size_text}{market} {month_label(contract)}")
    return "; ".join(legs)


def segment_pairs(directory: pd.DataFrame, wide: pd.DataFrame) -> pd.DataFrame:
    meta = directory.set_index("segment_id").to_dict("index")
    rows = []
    for long_id, short_id in permutations(directory.segment_id, 2):
        aligned = pd.concat([wide[long_id], wide[short_id]], axis=1).dropna()
        if len(aligned) < 31:
            continue
        long_s, short_s = aligned.iloc[:, 0], aligned.iloc[:, 1]
        relative = long_s - short_s
        st = series_stats(relative)
        diag = beta_diagnostics(long_s, short_s, pnl_sign=1)
        beta = diag["beta30_long_on_short"]
        lm, sm = meta[long_id], meta[short_id]
        rows.append(
            {
                "long_segment": long_id,
                "long_definition": lm["definition"],
                "short_segment": short_id,
                "short_definition": sm["definition"],
                "pair_type": (
                    "matched cross-market segment"
                    if lm["segment"] == sm["segment"] and lm["market"] != sm["market"]
                    else "same-market segment pair"
                    if lm["market"] == sm["market"]
                    else "mixed cross-market segment pair"
                ),
                "relative_level_long_minus_short_bp": st["current"],
                "relative_change_1d_bp": st["change_1d"],
                "trade_pnl_1to1_1d_bp": st["change_1d"],
                "trade_pnl_1to1_5d_bp": st["change_5d"],
                "relative_mean_bp": st["mean"],
                "relative_zscore": st["zscore"],
                "relative_percentile": st["percentile"],
                "relative_min_bp": st["minimum"],
                "relative_max_bp": st["maximum"],
                "entry_read": (
                    "cheap for long-segment/short-segment mean reversion"
                    if st["percentile"] <= 5
                    else "rich; inverse row is the mean-reversion direction"
                    if st["percentile"] >= 95
                    else "not a tail valuation"
                ),
                "legs_1to1": trade_legs(long_id, short_id, 1.0),
                "legs_beta_hedged": trade_legs(long_id, short_id, beta),
                **diag,
            }
        )
    return pd.DataFrame(rows)


def write_df(ws, df: pd.DataFrame):
    for col, name in enumerate(df.columns, 1):
        cell = ws.cell(1, col, name)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(color=WHITE, bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = Border(bottom=THIN)
    for row, values in enumerate(df.itertuples(index=False, name=None), 2):
        for col, value in enumerate(values, 1):
            if pd.isna(value):
                value = None
            elif isinstance(value, pd.Timestamp):
                value = value.to_pydatetime()
            elif isinstance(value, np.generic):
                value = value.item()
            cell = ws.cell(row, col, value)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = Border(bottom=THIN)
            if row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=PALE)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(df.columns))}{len(df)+1}"
    ws.row_dimensions[1].height = 32
    for i, name in enumerate(df.columns, 1):
        width = min(
            max([len(str(name))] + [len(str(v)) for v in df[name].head(100).dropna()]) + 2,
            42,
        )
        if any(x in name for x in ["definition", "legs_", "interpretation", "entry_read"]):
            width = 42
        ws.column_dimensions[get_column_letter(i)].width = width
        if any(x in name for x in ["beta", "corr", "vol", "zscore", "percentile", "_bp"]):
            for row in range(2, len(df) + 2):
                ws.cell(row, i).number_format = "0.00"
    for metric in ["spread_percentile", "relative_percentile", "percentile"]:
        if metric in df:
            col = get_column_letter(df.columns.get_loc(metric) + 1)
            ws.conditional_formatting.add(
                f"{col}2:{col}{len(df)+1}",
                ColorScaleRule(
                    start_type="num", start_value=0, start_color="63BE7B",
                    mid_type="num", mid_value=50, mid_color="FFEB84",
                    end_type="num", end_value=100, end_color="F8696B",
                ),
            )


def add_readme(wb: Workbook, asof: pd.Timestamp, n_contracts: int):
    ws = wb.active
    ws.title = "README"
    ws.sheet_view.showGridLines = False
    ws["A1"] = "Three-curve STIR exhaustive RV universe"
    ws["A1"].font = Font(size=18, bold=True, color=WHITE)
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws.merge_cells("A1:F1")
    rows = [
        ("As of", str(asof.date())),
        ("Sample", f"First common observations from {SAMPLE_START.date()} through as-of"),
        ("Contracts", f"{n_contracts} live fixed contracts; every ordered pair is included"),
        ("Market data", "Barchart fixed-contract EOD via headed browser session"),
        ("Canonical panel", "data/cache/stir_curves/panel.csv"),
        ("Auditable raw", "data/raw/stir_futures/histories_raw.json.gz"),
        ("Policy spots", "FRED local cache; context only, not pair beta inputs"),
        ("Contract-pair beta", "β30 = cov(Δ long rate, Δ short rate) / var(Δ short rate)"),
        (
            "Contract trade convention",
            "Long leg = buy futures/receive rate; short leg = sell futures/pay rate. "
            "P&L = −Δlong rate + hedge×Δshort rate.",
        ),
        (
            "Segment convention",
            "Every segment is back minus front. Long segment P&L = +Δsegment; "
            "short segment P&L = −Δsegment.",
        ),
        (
            "β interpretation",
            "β<1 means one short unit overhedges the long leg; β>1 means one short "
            "unit underhedges it. β is descriptive and can change regime.",
        ),
        (
            "Raw data sheets",
            "Raw_Rates, Raw_Prices, Raw_Changes and Raw_Segments contain every "
            "input used in valuation and beta calculations.",
        ),
    ]
    for row, (key, value) in enumerate(rows, 3):
        ws.cell(row, 1, key).font = Font(bold=True, color=NAVY)
        ws.cell(row, 2, value)
        ws.cell(row, 2).alignment = Alignment(wrap_text=True)
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=6)
        ws.row_dimensions[row].height = 32
    ws.column_dimensions["A"].width = 27
    for col in "BCDEF":
        ws.column_dimensions[col].width = 23


def curated_opportunities(
    pairs: pd.DataFrame,
    segments: pd.DataFrame,
    seg_pairs: pd.DataFrame,
) -> pd.DataFrame:
    contract_lookup = pairs.set_index(
        ["long_receive_leg", "short_pay_leg"]
    ).to_dict("index")
    segment_lookup = segments.set_index("segment_id").to_dict("index")
    segment_pair_lookup = seg_pairs.set_index(
        ["long_segment", "short_segment"]
    ).to_dict("index")

    def outright(
        market: str,
        segment: str,
        back: str,
        front: str,
        assessment: str,
        risk: str,
    ) -> dict:
        seg = segment_lookup[f"{market}_{segment}"]
        pair = contract_lookup[(f"{market}_{back}", f"{market}_{front}")]
        return {
            "trade": f"{market} {SEGMENTS[segment]['label']} flattener",
            "implementation": (
                f"PAY {market} {month_label(front)}; "
                f"RECEIVE {market} {month_label(back)}"
            ),
            "current_level_bp": seg["current_bp"],
            "percentile": seg["percentile"],
            "beta30_back_on_front": pair["beta30_long_on_short"],
            "corr30": pair["corr30"],
            "assessment": assessment,
            "main_risk": risk,
        }

    relative = segment_pair_lookup[("EURIBOR_JUN_CAL", "SONIA_JUN_CAL")]
    rows = [
        {
            "rank": 1,
            **outright(
                "EURIBOR",
                "JUN_CAL",
                "M28",
                "M27",
                "Best outright flattener entry after weekly re-steepening",
                "Latest selloffs steepened this gap; this is a contrarian restart call",
            ),
        },
        {
            "rank": 2,
            "trade": "Long Euribor June calendar vs short SONIA June calendar",
            "implementation": relative["legs_1to1"],
            "current_level_bp": relative["relative_level_long_minus_short_bp"],
            "percentile": relative["relative_percentile"],
            "beta30_back_on_front": relative["beta30_long_on_short"],
            "corr30": relative["corr30"],
            "assessment": "Best cross-market curve RV; hedges common curve shocks",
            "main_risk": "Relative level improved but is no longer beyond the 95th percentile",
        },
        {
            "rank": 3,
            **outright(
                "SOFR",
                "JUN_CAL",
                "M28",
                "M27",
                "Good statistical flattener entry; less direct macro support",
                "US curve can keep bear-steepening if long-forward inflation risk dominates",
            ),
        },
        {
            "rank": 4,
            **outright(
                "SONIA",
                "JUN_CAL",
                "M28",
                "M27",
                "Strong valuation but overlaps the live BoE peak-fade position",
                "Adds SONIA curve exposure while live beta remains elevated",
            ),
        },
        {
            "rank": 5,
            **outright(
                "SONIA",
                "REVERSAL",
                "Z28",
                "U27",
                "Fair flattener/hedge entry after steepening toward zero",
                "At equal size it cancels the Sep-27 leg of the live buildup trade",
            ),
        },
        {
            "rank": 6,
            **outright(
                "EURIBOR",
                "BUILDUP",
                "U27",
                "Z26",
                "Very rich statistically, but high-momentum and exposed to ECB hawkish tails",
                "Catching a peak-led selloff; scale only if fading 2027 policy pricing",
            ),
        },
    ]
    return pd.DataFrame(rows)


def build(asof: str | None = None) -> Path:
    panel, active, rates, prices, end = load_panel(asof)
    contracts = contract_directory(active, rates, prices)
    pairs = contract_pairs(contracts, rates)
    segments, segment_wide = segment_data(rates)
    seg_pairs = segment_pairs(segments, segment_wide)

    within = pairs[pairs.pair_type == "within market"].copy()
    cross = pairs[pairs.pair_type.str.startswith("cross market")].copy()
    matched_contracts = pairs[pairs.pair_type == "cross market, same month"].copy()
    matched_segments = seg_pairs[
        seg_pairs.pair_type == "matched cross-market segment"
    ].copy()

    contract_opps = pairs[
        (pairs.spread_percentile >= 95) & (pairs.corr30.abs() >= 0.65)
    ].sort_values(
        ["spread_percentile", "r_squared30"], ascending=[False, False]
    ).head(40)
    segment_opps = seg_pairs[
        (seg_pairs.relative_percentile <= 5) & (seg_pairs.corr30.abs() >= 0.60)
    ].sort_values(
        ["relative_percentile", "r_squared30"], ascending=[True, False]
    ).head(40)
    curated = curated_opportunities(pairs, segments, seg_pairs)

    raw_rates = rates.reset_index()
    raw_prices = prices.reset_index()
    raw_changes = rates.diff().reset_index()
    raw_segments = segment_wide.reset_index()

    wb = Workbook()
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    add_readme(wb, end, len(contracts))
    sheets = [
        ("Curated_Opportunities", curated),
        ("Contract_Directory", contracts),
        ("Segment_Directory", segments),
        ("Contract_Opportunities", contract_opps),
        ("Segment_Opportunities", segment_opps),
        ("Contract_Pairs_All", pairs),
        ("Within_Market_Pairs", within),
        ("Cross_Market_Pairs", cross),
        ("Matched_Bellies", matched_contracts),
        ("Segment_Pairs_All", seg_pairs),
        ("Matched_Segment_RV", matched_segments),
        ("Raw_Rates", raw_rates),
        ("Raw_Prices", raw_prices),
        ("Raw_Changes", raw_changes),
        ("Raw_Segments", raw_segments),
    ]
    for name, frame in sheets:
        ws = wb.create_sheet(name)
        write_df(ws, frame)
        ws.sheet_view.zoomScale = 85

    output = OUT_DIR / f"stir_full_rv_universe_{end.date()}.xlsx"
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    if asof is None:
        shutil.copyfile(output, LATEST)
    return output


def validate(path: Path):
    wb = load_workbook(path, read_only=True)
    required = {
        "README", "Contract_Directory", "Segment_Directory",
        "Contract_Pairs_All", "Segment_Pairs_All", "Raw_Rates", "Raw_Segments",
    }
    if missing := required - set(wb.sheetnames):
        raise RuntimeError(f"missing sheets: {sorted(missing)}")
    if wb["Contract_Pairs_All"].max_row < 1000:
        raise RuntimeError("contract pair universe unexpectedly small")
    if wb["Segment_Pairs_All"].max_row < 500:
        raise RuntimeError("segment pair universe unexpectedly small")
    wb.close()


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--asof")
    args = parser.parse_args()
    path = build(args.asof)
    validate(path)
    print(f"wrote {path} ({path.stat().st_size:,} bytes)")
    if args.asof is None:
        validate(LATEST)
        print(f"wrote {LATEST} ({LATEST.stat().st_size:,} bytes)")


if __name__ == "__main__":
    main()
